"""Harness optimization on Harvey's Legal Agent Benchmark with a NATIVE dspy.Flex + dspy.GEPA loop.

Inspired by Joel Niklaus, "Don't Train the Model, Evolve the Harness"
(https://huggingface.co/spaces/joelniklaus/harness-optimization): freeze the model, and let an LLM
proposer evolve the *harness* — the runtime code around the model — against Harvey's Legal Agent
Benchmark (LAB, github.com/harveyai/harvey-labs). The article's findings: off-the-shelf agent harnesses
scored BELOW the vanilla baseline, the failures were delivery/format not reasoning, and the biggest
gains came from deterministic CODE mechanisms — 5 of the top 6 were code, not prompts.

dspy.Flex + dspy.GEPA is the DSPy realization of exactly that: a Flex module's code (`module_src`) IS a
harness, and GEPA rewrites that code (not just prompts). This demo is deliberately NATIVE — no seeding,
no hand-written starter:

    harness = dspy.Flex(LegalDeliverable)          # starts as a single-call dspy.Predict harness
    optimized = dspy.GEPA(metric=rubric_score, ...).compile(harness, trainset=..., valset=...)

GEPA evolves that harness code (decompose the work, structure the deliverable, ground claims in the
documents, address each criterion) with the executor model frozen. The proposer is Opus (mirroring the
article's Opus 4.8); the frozen model is a cheap executor (Haiku).

USES THE REAL DATASET. Point HARVEY_LABS_DIR at a local clone; the demo loads real tasks (instructions,
documents, expert rubric) and scores with a faithful replica of LAB's rubric judge (per-criterion
PASS/FAIL vs `match_criteria`, default judge claude-sonnet-4-6). Metrics mirror LAB: pooled criterion
pass rate (dense — what GEPA optimizes) and all-pass rate (sparse headline).

    git clone https://github.com/harveyai/harvey-labs
    export HARVEY_LABS_DIR=/path/to/harvey-labs
    .venv/bin/python -m pytest tests/flex/demo/test_flex_harness_optimization.py -s

Honest adaptations (so a single Flex module can stand in for LAB's multi-file agent, cheaply):
  * LAB's harness writes multiple files to a sandbox with six tools; here the harness is one dspy.Flex
    producing the deliverable as TEXT, and every criterion is judged against that text. This exercises
    the content/structure harness, not the filesystem "landing to exact filenames" (which — ironically
    — was the article's #1 mechanism; it needs a real file sandbox).
  * The judge replicates LAB's methodology but BATCHES all of a task's criteria into one call (LAB
    judges each independently) — a big cost saving; it is not LAB's exact judge code.
  * Documents are extracted best-effort (needs python-docx / pypdf / openpyxl for Office/PDF; text
    formats always work) and truncated to bound tokens.

Cost is dominated by the harness's LLM calls over the (long) documents, so the defaults are tiny — a small
trainset, capped criteria, capped document length. Everything scales via env. Don't expect the article's
numbers at this scale; the point is the native loop and whether GEPA moves the pooled rate. Results are
reported, not asserted. Skips if HARVEY_LABS_DIR is unset or has no tasks.
"""

from __future__ import annotations

import json
import os
import random
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from dotenv import load_dotenv

import dspy
from dspy.teleprompt.gepa.gepa_utils import ScoreWithFeedback

load_dotenv()

LABS_DIR = os.getenv("HARVEY_LABS_DIR")
if not LABS_DIR or not (Path(LABS_DIR) / "tasks").is_dir():
    pytest.skip(
        "Set HARVEY_LABS_DIR to a local clone of github.com/harveyai/harvey-labs "
        "(the 'tasks/' directory must exist).",
        allow_module_level=True,
    )
TASKS_ROOT = Path(LABS_DIR) / "tasks"
DEMO_DIR = Path(__file__).parent
SRC_PATH = DEMO_DIR / "harvey_lab_harness.json"

# The frozen model inside the harness, the harness evolver (proposer), and LAB's rubric judge (its
# default is claude-sonnet-4-6). Override any via env.
EXEC_LM = dspy.LM(os.getenv("HARVEY_EXEC_LM", "anthropic/claude-haiku-4-5"), max_tokens=8000)
REFLECTION_LM = dspy.LM(os.getenv("HARVEY_REFLECTION_LM", "anthropic/claude-opus-4-8"), temperature=1.0, max_tokens=8000)
JUDGE_LM = dspy.LM(os.getenv("HARVEY_JUDGE_LM", "anthropic/claude-sonnet-4-6"), max_tokens=4000)

# Small defaults: cost is ~ (train+val eval rollouts + GEPA budget) harness forwards over the (long)
# documents. Keep these tiny for a first run; scale via env once you've seen the cost.
N_TRAIN = int(os.getenv("HARVEY_N_TRAIN", "3"))
N_VAL = int(os.getenv("HARVEY_N_VAL", "2"))
N_TEST = int(os.getenv("HARVEY_N_TEST", "3"))
MAX_CRITERIA = int(os.getenv("HARVEY_MAX_CRITERIA", "12"))  # cap criteria/task judged (tasks can have 50+)
MAX_DOC_CHARS = int(os.getenv("HARVEY_MAX_DOC_CHARS", "12000"))  # cap concatenated documents fed to the harness
GEPA_BUDGET = int(os.getenv("HARVEY_GEPA_BUDGET", "30"))  # must exceed one train+val eval to leave room for reflection
REFLECTION_MINIBATCH = 2
EVAL_THREADS = 4

_start = time.monotonic()


def _log(msg: str) -> None:
    """Timestamped, flushed progress line (visible live under `pytest -s`)."""
    print(f"[+{time.monotonic() - _start:5.0f}s] {msg}", flush=True)


class LegalDeliverable(dspy.Signature):
    """Produce the deliverable for a legal task so it satisfies an expert rubric.

    You are given the `task_instructions` and the source `documents`. Return `deliverable`: the full
    text of the work product (memo, analysis, draft, or redline) that carries out the instructions and
    would pass expert review — grounded in the documents, with the required facts, conclusions,
    citations, and structure. Do not restate the instructions; produce the work product itself.
    """

    task_instructions: str = dspy.InputField(desc="What the agent must do.")
    documents: str = dspy.InputField(desc="Concatenated source documents (may be truncated).")
    deliverable: str = dspy.OutputField(desc="The full work-product text.")


class RubricJudge(dspy.Signature):
    """Grade a legal deliverable against expert rubric criteria, in the style of Harvey LAB's judge.

    For each numbered criterion, decide PASS or FAIL strictly by its stated standard, judging ONLY the
    deliverable text (no benefit of the doubt for work that isn't present). Return `verdicts` as a list
    of "PASS"/"FAIL", one per criterion, in the same order as the rubric.
    """

    task_title: str = dspy.InputField()
    deliverable: str = dspy.InputField(desc="The work product to grade.")
    criteria: str = dspy.InputField(desc="Numbered rubric criteria; each states its own PASS/FAIL standard.")
    verdicts: list[str] = dspy.OutputField(desc='One "PASS" or "FAIL" per criterion, in order.')


# --- data: load real LAB tasks -----------------------------------------------


def _extract_text(path: Path, budget: int) -> str:
    """Best-effort text from a document file, up to ~`budget` chars. Office/PDF need optional libs;
    text formats always work; anything else is skipped."""
    if budget <= 0:
        return ""
    suffix = path.suffix.lower()
    try:
        if suffix in (".txt", ".md", ".csv", ".json", ".html", ""):
            return path.read_text(encoding="utf-8", errors="ignore")[:budget]
        if suffix == ".docx":
            import docx  # python-docx

            return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)[:budget]
        if suffix == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(str(path))
            return "\n".join((page.extract_text() or "") for page in reader.pages)[:budget]
        if suffix in (".xlsx", ".xlsm"):
            import openpyxl

            book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows = [
                "\t".join("" if cell is None else str(cell) for cell in row)
                for sheet in book.worksheets
                for row in sheet.iter_rows(values_only=True)
            ]
            return "\n".join(rows)[:budget]
    except Exception:
        return ""
    return ""


def _load_documents(task_dir: Path, budget: int) -> str:
    docs_dir = task_dir / "documents"
    if not docs_dir.is_dir():
        return ""
    parts, used = [], 0
    for file in sorted(p for p in docs_dir.rglob("*") if p.is_file()):
        if used >= budget:
            break
        text = _extract_text(file, budget - used).strip()
        if text:
            chunk = f"===== {file.name} =====\n{text}"
            parts.append(chunk)
            used += len(chunk)
    return "\n\n".join(parts)


def _load_task(task_dir: Path) -> dict:
    data = json.loads((task_dir / "task.json").read_text(encoding="utf-8"))
    criteria = [
        {"id": c.get("id", ""), "title": c.get("title", ""), "match": c.get("match_criteria", "")}
        for c in data.get("criteria", [])
        if c.get("match_criteria")
    ][:MAX_CRITERIA]
    return {
        "task_id": str(task_dir.relative_to(TASKS_ROOT)),
        "title": data.get("title", ""),
        "task_instructions": data.get("instructions", ""),  # LAB's json field is "instructions"
        "criteria": criteria,
        "documents": _load_documents(task_dir, MAX_DOC_CHARS),
    }


def _splits() -> tuple[list, list, list]:
    """Load a small deterministic sample of real tasks. HARVEY_TASKS ('area/task,area/task') overrides
    the sample; otherwise tasks are drawn at random (seeded) from the whole benchmark."""
    override = os.getenv("HARVEY_TASKS")
    if override:
        chosen = [TASKS_ROOT / t.strip() for t in override.split(",") if t.strip()]
    else:
        all_dirs = sorted(p.parent for p in TASKS_ROOT.rglob("task.json"))
        rng = random.Random(0)
        rng.shuffle(all_dirs)
        chosen = all_dirs[: N_TRAIN + N_VAL + N_TEST]
    examples = [
        dspy.Example(**_load_task(d)).with_inputs("task_instructions", "documents")
        for d in chosen
        if (d / "task.json").is_file()
    ]
    examples = [e for e in examples if e.criteria]  # can't score a task with no criteria
    return examples[:N_TRAIN], examples[N_TRAIN : N_TRAIN + N_VAL], examples[N_TRAIN + N_VAL :]


# --- rubric scoring (a faithful, batched replica of LAB's judge) -------------


def _judge(title: str, deliverable: str, criteria: list[dict]) -> list[bool]:
    """PASS/FAIL per criterion for one deliverable, in a single batched judge call. Missing/short
    verdicts count as FAIL (a delivery that can't be graded doesn't pass)."""
    if not criteria:
        return []
    rubric = "\n".join(f"{i + 1}. [{c['title']}] {c['match']}" for i, c in enumerate(criteria))
    try:
        with dspy.context(lm=JUDGE_LM):
            out = dspy.Predict(RubricJudge)(task_title=title, deliverable=deliverable or "(empty)", criteria=rubric)
        verdicts = [str(v).strip().upper().startswith("PASS") for v in out.verdicts]
    except Exception:
        verdicts = []
    return (verdicts + [False] * len(criteria))[: len(criteria)]


def _score(title: str, deliverable: str, criteria: list[dict]) -> tuple[float, bool]:
    """(pooled criterion pass rate, all-pass) — LAB's two metrics."""
    verdicts = _judge(title, deliverable, criteria)
    passed = sum(verdicts)
    return passed / len(criteria), passed == len(criteria)


def _metric(gold, pred, trace=None, pred_name=None, pred_trace=None) -> ScoreWithFeedback:
    """GEPA optimizes the pooled criterion rate; feedback names the failing criteria so the proposer
    can improve the harness (structure, grounding, coverage) rather than the model."""
    verdicts = _judge(gold.title, getattr(pred, "deliverable", "") or "", gold.criteria)
    passed, n = sum(verdicts), len(gold.criteria)
    fails = [gold.criteria[i]["title"] for i, ok in enumerate(verdicts) if not ok][:8]
    if not fails:
        feedback = f"All {n} rubric criteria passed. Keep the harness structured and grounded in the documents."
    else:
        feedback = (
            f"{passed}/{n} rubric criteria passed. Failing: " + "; ".join(fails) + ". Improve the HARNESS "
            "(the module's code), not the model: decompose the work, read/ground every claim in the "
            "documents, address each criterion explicitly, and produce a complete, well-structured deliverable."
        )
    return ScoreWithFeedback(score=passed / n if n else 0.0, feedback=feedback)


def _evaluate(program: dspy.Module, dataset: list) -> tuple[float, float]:
    """(mean pooled criterion rate, all-pass rate) over a dataset."""

    def run_one(ex):
        try:
            with dspy.context(lm=EXEC_LM):
                pred = program(task_instructions=ex.task_instructions, documents=ex.documents)
            return _score(ex.title, getattr(pred, "deliverable", "") or "", ex.criteria)
        except Exception:
            return 0.0, False

    with ThreadPoolExecutor(max_workers=EVAL_THREADS) as pool:
        rows = list(pool.map(run_one, dataset))
    pooled = sum(p for p, _ in rows) / len(rows)
    all_pass = sum(a for _, a in rows) / len(rows)
    return pooled, all_pass


def test_flex_harvey_lab_harness() -> None:
    global _start
    _start = time.monotonic()
    dspy.configure(lm=EXEC_LM)
    train, val, test = _splits()
    if not (train and test):
        pytest.skip("Not enough scorable tasks loaded — check HARVEY_LABS_DIR / HARVEY_TASKS / N_* settings.")
    _log(f"Harvey LAB harness | train={len(train)} val={len(val)} test={len(test)} tasks "
         f"(<= {MAX_CRITERIA} criteria/task, docs <= {MAX_DOC_CHARS} chars)")
    for e in train + val + test:
        _log(f"  {e.task_id}: {len(e.criteria)} criteria, {len(e.documents)} doc chars")

    # 1. Native baseline: dspy.Flex starts as a single-call dspy.Predict harness. No seeding.
    harness = dspy.Flex(LegalDeliverable)
    assert "dspy.Predict(" in (harness.module_src or ""), "expected the native Predict baseline harness"
    base_pooled, base_allpass = _evaluate(harness, test)
    _log(f"[baseline Predict harness] pooled={base_pooled:.1%}  all-pass={base_allpass:.1%}")

    # 2. GEPA evolves the harness CODE (module_src); the executor model stays frozen.
    _log("evolving the harness with GEPA (Opus proposes; Haiku frozen)...")
    optimized = dspy.GEPA(
        metric=_metric,
        reflection_lm=REFLECTION_LM,
        max_metric_calls=GEPA_BUDGET,
        reflection_minibatch_size=REFLECTION_MINIBATCH,
        num_threads=EVAL_THREADS,
        track_stats=True,
    ).compile(harness, trainset=train, valset=val)

    opt_pooled, opt_allpass = _evaluate(optimized, test)
    _log(f"[evolved harness] pooled={opt_pooled:.1%}  all-pass={opt_allpass:.1%}")
    _log(f"pooled {base_pooled:.1%} -> {opt_pooled:.1%} ({opt_pooled - base_pooled:+.1%}); "
         f"harness code changed: {optimized.module_src != harness.module_src}")

    SRC_PATH.write_text(
        json.dumps(
            {
                "baseline_module_src": harness.module_src,
                "evolved_module_src": optimized.module_src,
                "baseline": {"pooled": base_pooled, "all_pass": base_allpass},
                "optimized": {"pooled": opt_pooled, "all_pass": opt_allpass},
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    _log(f"saved baseline + evolved harness source and scores -> {SRC_PATH}")
    _log("done.")

    # Invariants only: whether GEPA moves the pooled rate depends on the live models/budget, so it's
    # reported rather than asserted.
    assert optimized.module_src is not None
    assert all(0.0 <= r <= 1.0 for r in (base_pooled, base_allpass, opt_pooled, opt_allpass))


if __name__ == "__main__":
    test_flex_harvey_lab_harness()
